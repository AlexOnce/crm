import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # Configurar formato 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Paleta de Colores TAMep
    NAVY = RGBColor(15, 23, 42)       # #0f172a
    DARK_BLUE = RGBColor(14, 116, 144) # #0e7490
    CYAN = RGBColor(2, 132, 199)       # #0284c7
    EMERALD = RGBColor(5, 150, 105)    # #059669
    WHITE = RGBColor(255, 255, 255)
    SLATE = RGBColor(100, 116, 139)    # #64748b
    LIGHT_BG = RGBColor(248, 250, 252) # #f8fafc
    BORDER_COLOR = RGBColor(226, 232, 240)

    slides_data = [
        {
            "num": 1,
            "is_cover": True,
            "title": "TAMep WhatsApp CRM",
            "subtitle": "Manual de Operaciones y Ventas para Asesores\nTransporte Aéreo Militar - Empresa Pública",
            "notes": "Bienvenidos al entrenamiento del CRM oficial de WhatsApp de TAMep. Esta plataforma centraliza la atención de nuestros pasajeros, automatiza las consultas frecuentes mediante un bot inteligente y permite a nuestros asesores de ventas brindar una atención humana cálida, rápida y emitir boletos aéreos de forma segura en tiempo real."
        },
        {
            "num": 2,
            "title": "1. Acceso a la Plataforma (Inicio de Sesión)",
            "subtitle": "Ingreso seguro al entorno de trabajo con credenciales individuales",
            "points": [
                "Acceso vía navegador web en puesto de trabajo (URL oficial del CRM).",
                "Ingreso con correo corporativo asignado (ej. agente@tamep.bo).",
                "Contraseña individual y privada (nunca compartir credenciales).",
                "Protección de cuenta con limitador contra intentos repetidos de acceso.",
                "Sesión única por asesor para trazabilidad de ventas y comisiones."
            ],
            "capture_hint": "Pegar aquí Captura Real de la Pantalla de Login (http://127.0.0.1:8000/login)",
            "notes": "Cada asesor cuenta con credenciales únicas. Nunca compartan su contraseña con otros compañeros, ya que todas las atenciones, mensajes enviados y boletos vendidos quedan registrados bajo su usuario."
        },
        {
            "num": 3,
            "title": "2. Perfil del Asesor y Seguridad de la Cuenta",
            "subtitle": "Autogestión de contraseña, descanso visual y preferencias",
            "points": [
                "Acceso directo: Hacer clic sobre el nombre del agente en la esquina superior derecha.",
                "Pestaña 'Información y Preferencias': Alternar Modo Claro / Modo Oscuro.",
                "Alertas y Sonidos: Configuración de notificaciones ante nuevos clientes.",
                "Pestaña 'Seguridad y Contraseña': Cambio de contraseña directo sin soporte técnico.",
                "Validación de seguridad: Requiere contraseña actual y nueva clave (mínimo 8 caracteres)."
            ],
            "capture_hint": "Pegar aquí Captura Real del Modal 'Perfil del Asesor' (Pestaña Información y Cambio de Contraseña)",
            "notes": "La barra superior se ha limpiado de botones innecesarios. Ahora, al presionar su nombre, acceden directamente a su perfil. En la primera sesión, recomendamos cambiar la contraseña temporal por una segura en la pestaña de Seguridad. También pueden activar el modo oscuro si trabajan en turnos nocturnos."
        },
        {
            "num": 4,
            "title": "3. Organización de la Bandeja de Entrada",
            "subtitle": "Estructura de chats para priorizar y gestionar la atención",
            "points": [
                "Pestaña 'En Espera' (🛎️): Pasajeros que solicitaron un asesor humano tras interactuar con el bot.",
                "Contador de alerta: Muestra en tiempo real la cantidad de clientes esperando atención.",
                "Pestaña 'Mis Chats' (💬): Conversaciones que tú tienes asignadas activamente en este momento.",
                "Pestaña 'Todos': Vista general y monitoreo de chats del día.",
                "Buscador integrado: Búsqueda instantánea por nombre de cliente o número celular."
            ],
            "capture_hint": "Pegar aquí Captura Real de la Barra Lateral Izquierda (Pestañas 'Mis Chats' y 'En Espera')",
            "notes": "Nuestra prioridad siempre es la pestaña 'En Espera'. Cada vez que un pasajero solicita un asesor, el sistema emite una alerta sonora y visual. Los agentes disponibles deben tomar estos chats para mantener nuestro tiempo de respuesta en menos de 2 minutos."
        },
        {
            "num": 5,
            "title": "4. Cómo Tomar una Conversación",
            "subtitle": "Paso de la atención automática (Bot) a la atención personalizada (Asesor)",
            "points": [
                "Paso 1: Seleccionar el chat de la lista 'En Espera'.",
                "Paso 2: Revisar el historial previo (ruta, fecha y consultas hechas al bot).",
                "Paso 3: Presionar el botón azul 'Tomar Chat' en la cabecera.",
                "Paso 4: El bot se pausa automáticamente y el chat pasa a tu pestaña 'Mis Chats'.",
                "Protección en tiempo real: Si otro asesor lo toma primero, se actualiza al instante."
            ],
            "capture_hint": "Pegar aquí Captura Real de la Cabecera del Chat con el Botón Azul 'Tomar Chat'",
            "notes": "Al hacer clic en 'Tomar Chat', el sistema te asigna la conversación de forma exclusiva en tiempo real. Los demás asesores verán que el chat ya fue atendido, evitando que dos personas respondan a la vez al mismo cliente."
        },
        {
            "num": 6,
            "title": "5. Interacción y Envío de Archivos Multimedia",
            "subtitle": "Canal oficial de WhatsApp para comunicación fluida con el pasajero",
            "points": [
                "Mensajes de Texto: Escritura fluida. Salto de línea con Shift + Enter y envío con Enter.",
                "Envío de Archivos (📎): Soporta imágenes (comprobantes de pago QR o transferencias).",
                "Documentos (PDF): Envío de cotizaciones oficiales o políticas de equipaje.",
                "Notas de Voz: Recepción y reproducción de mensajes de voz del pasajero.",
                "Navegación por Fecha: Botón 'Fecha' para saltar a Hoy, Ayer o fechas anteriores."
            ],
            "capture_hint": "Pegar aquí Captura Real del Área de Chat (Burbujas de Mensajes, Botón Clip y Redactor)",
            "notes": "A través de este canal pueden recibir los comprobantes de pago por transferencia bancaria o QR de los pasajeros, así como fotos de cédulas de identidad o pasaportes para la emisión."
        },
        {
            "num": 7,
            "title": "6. Panel Lateral de Reserva y Consulta de Vuelos",
            "subtitle": "Carga y verificación de datos del pasajero en tiempo real",
            "points": [
                "Datos pre-cargados por el Bot: Nombre, Cédula de Identidad, Ruta y Fecha de Vuelo.",
                "Edición por el asesor: Corrección de nombres, número de documento o ruta si el cliente cambia de plan.",
                "Consulta de Itinerarios TAMep: Verificación de disponibilidad y tarifas oficiales.",
                "Pasajeros múltiples: Configuración de cantidad de pasajeros con cálculo automático del precio total.",
                "Ficha completa: Vista clara de los datos antes de proceder a la venta."
            ],
            "capture_hint": "Pegar aquí Captura Real del Panel Lateral Derecho de Datos del Pasajero",
            "notes": "En el panel lateral derecho tienen la ficha completa del lead. Verifiquen siempre que el número de carnet y el nombre estén exactamente iguales al documento oficial para evitar inconvenientes en el aeropuerto."
        },
        {
            "num": 8,
            "title": "7. Emisión y Registro de Venta del Boleto",
            "subtitle": "Proceso oficial de cierre de venta y emisión transaccional",
            "points": [
                "Paso 1: Confirmar el pago y hacer clic en el botón verde 'Emitir Boleto'.",
                "Paso 2: Ventana de confirmación: '¿Está seguro que vendió el boleto?'.",
                "Paso 3: Al confirmar, el sistema genera el código de reserva PNR oficial (ej. TM-7052).",
                "Paso 4: Alerta de éxito: '¡Boleto Vendido! La venta se ha registrado exitosamente'.",
                "Paso 5: Traspaso limpio al bot: Finaliza la atención humana sin saturar el chat de mensajes."
            ],
            "capture_hint": "Pegar aquí Captura Real del Diálogo SweetAlert ('¿Está seguro que vendió el boleto?' / '¡Boleto Vendido!')",
            "notes": "Al emitir el boleto, el sistema procesa la transacción de manera limpia y segura. El PNR queda registrado en la base de datos oficial de TAMep, el contador del dashboard se actualiza y el pasajero vuelve a tener el bot activo para futuras consultas."
        },
        {
            "num": 9,
            "title": "8. Trabajo en Equipo y Atribución de Ventas",
            "subtitle": "Reglas de continuidad de chats y registro de comisiones",
            "points": [
                "¿Qué ocurre si un asesor toma un chat y otro lo continúa y vende?",
                "Reasignación automática: Al enviar una respuesta, el sistema asigna el chat al asesor activo.",
                "Atribución de la venta: El boleto y la venta se acreditan al asesor que emite el pasaje.",
                "Historial transparente: Cada mensaje conserva la autoría del agente que lo escribió.",
                "Métricas de Gerencia: Los boletos vendidos se reflejan en el perfil del agente en el Admin Dashboard."
            ],
            "capture_hint": "Pegar aquí Captura Real de un Chat con Mensajes de Agente o Vista de Métricas",
            "notes": "Si un compañero sale a refrigerio o termina su turno dejando un chat a medias, el asesor que toma el relevo y culmina la atención se queda con la venta registrada a su nombre en el panel de métricas de gerencia."
        },
        {
            "num": 10,
            "title": "9. Buenas Prácticas para una Atención de Excelencia",
            "subtitle": "Compromiso con la calidad del servicio aeronáutico de TAMep",
            "points": [
                "⏱️ Rapidez de Respuesta: Atender chats en espera dentro de los primeros 2 minutos.",
                "🔍 Precisión en Datos: Validar siempre ortografía del nombre y número de documento antes de emitir.",
                "🤖 Devolución al Bot: Si el cliente no compra o solo consulta horarios, usar 'Reactivar Bot'.",
                "🔒 Seguridad al salir: Cerrar sesión en el perfil al terminar el turno de trabajo.",
                "✈️ Compromiso TAMep: Calidez, respeto y vocación de servicio con todos nuestros pasajeros."
            ],
            "capture_hint": "Pegar aquí Captura Real de las Opciones de Cabecera ('Reactivar Bot' y Menú de Perfil)",
            "notes": "Siguiendo estas prácticas garantizamos una experiencia de 5 estrellas para nuestros pasajeros y aseguramos que cada una de sus ventas esté perfectamente registrada en el sistema. ¡Gracias por conectar a Bolivia con TAMep!"
        }
    ]

    blank_layout = prs.slide_layouts[6] # Blank slide

    for s_data in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        
        # Guardar notas del presentador
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = s_data["notes"]
        
        if s_data.get("is_cover"):
            # Diapositiva de Portada
            bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
            bg.fill.solid()
            bg.fill.fore_color.rgb = NAVY
            bg.line.fill.background()
            
            # Cinta de acento cyan
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1.8), Inches(0.15), Inches(3.8))
            accent.fill.solid()
            accent.fill.fore_color.rgb = CYAN
            accent.line.fill.background()
            
            # Cuadro de texto principal
            tx_box = slide.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.5), Inches(3.8))
            tf = tx_box.text_frame
            tf.word_wrap = True
            
            p1 = tf.paragraphs[0]
            p1.text = "TAMep"
            p1.font.bold = True
            p1.font.size = Pt(28)
            p1.font.color.rgb = CYAN
            
            p2 = tf.add_paragraph()
            p2.text = "CRM WHATSAPP"
            p2.font.bold = True
            p2.font.size = Pt(54)
            p2.font.color.rgb = WHITE
            
            p3 = tf.add_paragraph()
            p3.text = "Manual de Operaciones y Capacitación para Asesores de Ventas"
            p3.font.size = Pt(22)
            p3.font.color.rgb = RGBColor(186, 230, 253)
            p3.space_before = Pt(20)
            
            p4 = tf.add_paragraph()
            p4.text = "Transporte Aéreo Militar - Empresa Pública (Bolivia)"
            p4.font.size = Pt(14)
            p4.font.color.rgb = SLATE
            p4.space_before = Pt(30)
            
        else:
            # Fondo de diapositiva de contenido
            bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
            bg.fill.solid()
            bg.fill.fore_color.rgb = LIGHT_BG
            bg.line.fill.background()
            
            # Cabecera superior
            header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.3))
            header.fill.solid()
            header.fill.fore_color.rgb = NAVY
            header.line.fill.background()
            
            # Texto Cabecera
            h_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.15), Inches(11.7), Inches(1.0))
            h_tf = h_box.text_frame
            h_tf.word_wrap = True
            
            hp1 = h_tf.paragraphs[0]
            hp1.text = s_data["title"]
            hp1.font.bold = True
            hp1.font.size = Pt(26)
            hp1.font.color.rgb = WHITE
            
            hp2 = h_tf.add_paragraph()
            hp2.text = s_data["subtitle"]
            hp2.font.size = Pt(13)
            hp2.font.color.rgb = RGBColor(186, 230, 253)
            
            # Columna Izquierda: Tarjeta de Contenidos y Pasos
            left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.4))
            left_card.fill.solid()
            left_card.fill.fore_color.rgb = WHITE
            left_card.line.color.rgb = BORDER_COLOR
            left_card.line.width = Pt(1.5)
            
            # Texto Columna Izquierda
            c_box = slide.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.4), Inches(5.0))
            c_tf = c_box.text_frame
            c_tf.word_wrap = True
            
            cp_title = c_tf.paragraphs[0]
            cp_title.text = "Puntos Clave y Procedimiento:"
            cp_title.font.bold = True
            cp_title.font.size = Pt(17)
            cp_title.font.color.rgb = DARK_BLUE
            cp_title.space_after = Pt(14)
            
            for pt in s_data.get("points", []):
                p = c_tf.add_paragraph()
                p.text = "• " + pt
                p.font.size = Pt(13.5)
                p.font.color.rgb = RGBColor(30, 41, 59)
                p.space_after = Pt(10)
                
            # Columna Derecha: Tarjeta para Captura Real de Pantalla
            right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.4))
            right_card.fill.solid()
            right_card.fill.fore_color.rgb = RGBColor(241, 245, 249)
            right_card.line.color.rgb = RGBColor(203, 213, 225)
            right_card.line.width = Pt(1.5)
            
            # Recuadro interior de marcador de posición para imagen
            img_box = slide.shapes.add_textbox(Inches(7.3), Inches(2.2), Inches(5.0), Inches(4.2))
            img_tf = img_box.text_frame
            img_tf.word_wrap = True
            
            ip1 = img_tf.paragraphs[0]
            ip1.alignment = PP_ALIGN.CENTER
            ip1.text = "📸 ESPACIO PARA CAPTURA REAL"
            ip1.font.bold = True
            ip1.font.size = Pt(16)
            ip1.font.color.rgb = CYAN
            ip1.space_after = Pt(14)
            
            ip2 = img_tf.add_paragraph()
            ip2.alignment = PP_ALIGN.CENTER
            ip2.text = s_data["capture_hint"]
            ip2.font.size = Pt(13)
            ip2.font.color.rgb = SLATE
            ip2.space_after = Pt(18)
            
            ip3 = img_tf.add_paragraph()
            ip3.alignment = PP_ALIGN.CENTER
            ip3.text = "(Toma la captura en tu pantalla con Windows + Shift + S y pégala aquí)"
            ip3.font.italic = True
            ip3.font.size = Pt(11)
            ip3.font.color.rgb = RGBColor(148, 163, 184)

    output_path = os.path.join(os.path.dirname(__file__), "Manual_Agente_Ventas_TAMep.pptx")
    prs.save(output_path)
    print(f"Presentación generada exitosamente en: {output_path}")

if __name__ == "__main__":
    create_presentation()
